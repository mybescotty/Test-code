import os
import io
import json
import math
import re
import traceback
from datetime import datetime

import numpy as np
import pandas as pd
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 200 * 1024 * 1024  # 200 MB

# ─────────────────────────────────────────────────────────────────────────────
# Column-name normalisation helpers
# ─────────────────────────────────────────────────────────────────────────────

def normalise(s):
    """Lowercase, strip, collapse whitespace."""
    if not isinstance(s, str):
        s = str(s)
    return re.sub(r'\s+', ' ', s.strip().lower())


STOCK_ALIASES = {
    'part_number':  ['part number', 'part no', 'part#', 'partno', 'item number',
                     'item no', 'part code', 'stock code', 'part'],
    'description':  ['description', 'desc', 'part description', 'item description'],
    'qty_on_hand':  ['qty on hand', 'quantity on hand', 'qty oh', 'on hand', 'qoh',
                     'stock qty', 'balance', 'current stock', 'qty'],
    'unit_cost':    ['unit cost', 'cost', 'price', 'unit price', 'std cost'],
    'store':        ['store', 'location', 'storeroom', 'warehouse', 'store/location'],
}

TRANS_ALIASES = {
    'part_number':  ['part number', 'part no', 'part#', 'partno', 'item number',
                     'item no', 'part code', 'stock code', 'part'],
    'posting_date': ['posting date', 'date', 'transaction date', 'posted date',
                     'trans date', 'post date'],
    'trans_type':   ['transaction type', 'trans type', 'type', 'tran type',
                     'transaction', 'trans'],
    'trans_qty':    ['transaction qty', 'trans qty', 'quantity', 'qty',
                     'transaction quantity', 'trans quantity', 'amount'],
    'work_order':   ['work order', 'wo', 'work order no', 'wo number', 'work order #'],
    'asset':        ['fleet/bus number', 'asset', 'bus number', 'fleet number',
                     'equipment', 'asset number', 'fleet', 'bus no'],
    'to_from_code': ['to/from code', 'tocode', 'fromcode', 'to code', 'from code',
                     'to/from', 'store code', 'destination'],
}

MINMAX_ALIASES = {
    'part_number':  ['part number', 'part no', 'part#', 'partno', 'item number',
                     'item no', 'part code', 'stock code', 'part'],
    'current_min':  ['current min', 'min', 'minimum', 'min qty', 'min level'],
    'current_max':  ['current max', 'max', 'maximum', 'max qty', 'max level'],
    'lead_time':    ['lead time', 'lead time days', 'lt', 'lead days',
                     'replenishment time'],
}


def detect_columns(df, alias_map):
    """
    Returns dict mapping logical_key -> actual_column_name (or None).
    Uses normalised string matching.
    """
    norm_cols = {normalise(c): c for c in df.columns}
    mapping = {}
    for key, aliases in alias_map.items():
        found = None
        for alias in aliases:
            if normalise(alias) in norm_cols:
                found = norm_cols[normalise(alias)]
                break
        mapping[key] = found
    return mapping


def apply_mapping(df, col_map, user_overrides=None):
    """
    Renames df columns according to col_map (logical_key -> actual col).
    user_overrides can override auto-detected mappings.
    """
    if user_overrides:
        for key, col in user_overrides.items():
            if col and col in df.columns:
                col_map[key] = col
    rename = {v: k for k, v in col_map.items() if v is not None}
    return df.rename(columns=rename)


# ─────────────────────────────────────────────────────────────────────────────
# Core MRP analysis
# ─────────────────────────────────────────────────────────────────────────────

def classify_consumption(months_with_usage, total_months=12):
    if months_with_usage == 0:
        return 'Dead Stock'
    if months_with_usage >= total_months * 0.5:
        return 'Fast Mover'
    return 'Slow Mover'


def mode_order_qty(series):
    """Most common positive REC qty for a part; fallback 10."""
    vals = series[series > 0].round(0)
    if vals.empty:
        return 10
    counts = vals.value_counts()
    return int(counts.idxmax()) if not counts.empty else 10


def run_mrp_analysis(stock_df, trans_df, minmax_df,
                     stock_map, trans_map, minmax_map,
                     stock_overrides=None, trans_overrides=None, minmax_overrides=None):
    """
    Main analysis. Returns:
      recommendations_df, exceptions_df, monthly_df, summary_stats, chart_data
    """
    # ── Apply column mappings ──────────────────────────────────────────────
    stock  = apply_mapping(stock_df.copy(),  stock_map,  stock_overrides)
    trans  = apply_mapping(trans_df.copy(),  trans_map,  trans_overrides)
    minmax = apply_mapping(minmax_df.copy(), minmax_map, minmax_overrides)

    # ── Coerce types ──────────────────────────────────────────────────────
    for df, col in [(stock, 'part_number'), (trans, 'part_number'),
                    (minmax, 'part_number')]:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip().str.upper()

    # qty on hand
    if 'qty_on_hand' in stock.columns:
        stock['qty_on_hand'] = pd.to_numeric(stock['qty_on_hand'], errors='coerce').fillna(0)
    else:
        stock['qty_on_hand'] = 0

    if 'unit_cost' in stock.columns:
        stock['unit_cost'] = pd.to_numeric(stock['unit_cost'], errors='coerce').fillna(0)
    else:
        stock['unit_cost'] = 0

    # transaction qty
    if 'trans_qty' in trans.columns:
        trans['trans_qty'] = pd.to_numeric(trans['trans_qty'], errors='coerce').fillna(0)
    else:
        raise ValueError("Transaction quantity column could not be identified.")

    # posting date → month period
    if 'posting_date' in trans.columns:
        trans['posting_date'] = pd.to_datetime(trans['posting_date'], errors='coerce',
                                               dayfirst=True)
        trans['month'] = trans['posting_date'].dt.to_period('M')
    else:
        raise ValueError("Posting date column could not be identified.")

    # normalise text columns
    for col in ['trans_type', 'to_from_code']:
        if col in trans.columns:
            trans[col] = trans[col].astype(str).str.strip().str.upper()
        else:
            trans[col] = ''

    # min/max numerics
    for col in ['current_min', 'current_max']:
        if col in minmax.columns:
            minmax[col] = pd.to_numeric(minmax[col], errors='coerce').fillna(0)
        else:
            minmax[col] = 0

    if 'lead_time' in minmax.columns:
        minmax['lead_time'] = pd.to_numeric(minmax['lead_time'], errors='coerce')
    else:
        minmax['lead_time'] = np.nan

    # ── Filter to consumption transactions ────────────────────────────────
    #
    # CONSUMPTION = True for:
    #   I  + to_from_code=EVNT + negative qty  → issue to bus
    #   I  + to_from_code=STOR + positive qty  → return (negative consumption)
    #   REC + to_from_code=EVNT               → direct to bus
    #   STTK + negative qty                   → unrecorded usage
    #
    # NOT consumption:
    #   REC + to_from_code=STOR               → goods in
    #   STTK + positive qty                   → stock found

    def consumption_qty(row):
        t = row.get('trans_type', '')
        c = row.get('to_from_code', '')
        q = row.get('trans_qty', 0)

        if t == 'I':
            if c == 'EVNT' and q < 0:
                return abs(q)       # issue to bus — positive consumption
            if c == 'STOR' and q > 0:
                return -q           # return to stock — negative consumption
        elif t == 'REC':
            if c == 'EVNT':
                return abs(q)       # direct delivery to bus
            # REC to STOR = goods in, not consumption
        elif t == 'STTK':
            if q < 0:
                return abs(q)       # unrecorded usage
            # positive STTK = stock found, exclude
        return 0

    trans['consumption'] = trans.apply(consumption_qty, axis=1)

    # also tag REC-to-STOR for order qty mode calculation
    trans['is_goods_in'] = (
        (trans['trans_type'] == 'REC') & (trans['to_from_code'] == 'STOR')
    )

    # ── Monthly consumption per part ──────────────────────────────────────
    consumption = trans[trans['consumption'] != 0].copy()
    monthly = (
        consumption.groupby(['part_number', 'month'])['consumption']
        .sum()
        .reset_index()
    )

    # Build full part × month matrix
    all_parts = stock['part_number'].unique().tolist()
    # include parts in trans/minmax not in stock
    for df in [trans, minmax]:
        if 'part_number' in df.columns:
            for p in df['part_number'].unique():
                if p not in all_parts:
                    all_parts.append(p)

    all_months = sorted(trans['month'].dropna().unique())
    # limit to 12 most recent months
    all_months = all_months[-12:] if len(all_months) > 12 else all_months

    idx = pd.MultiIndex.from_product([all_parts, all_months],
                                      names=['part_number', 'month'])
    monthly_full = (
        monthly.set_index(['part_number', 'month'])
        .reindex(idx, fill_value=0)
        .reset_index()
    )
    monthly_full['month_str'] = monthly_full['month'].astype(str)

    # ── Per-part statistics ───────────────────────────────────────────────
    part_stats = []

    # Pre-index goods_in by part for mode calculation
    goods_in_by_part = (
        trans[trans['is_goods_in']]
        .groupby('part_number')['trans_qty']
    )

    minmax_idx = minmax.set_index('part_number') if 'part_number' in minmax.columns else pd.DataFrame()
    stock_idx  = stock.set_index('part_number')  if 'part_number' in stock.columns  else pd.DataFrame()

    for part in all_parts:
        part_monthly = monthly_full[monthly_full['part_number'] == part]['consumption'].values
        months_with_usage = int((part_monthly > 0).sum())
        total_consumption = float(part_monthly.sum())
        mean_monthly      = float(np.mean(part_monthly))
        std_monthly       = float(np.std(part_monthly, ddof=0))
        peak_monthly      = float(np.max(part_monthly))

        classification = classify_consumption(months_with_usage, len(all_months))

        # lead time
        lead_days = 7
        if part in minmax_idx.index and 'lead_time' in minmax_idx.columns:
            lt = minmax_idx.at[part, 'lead_time']
            if pd.notna(lt) and lt > 0:
                lead_days = float(lt)

        # current min/max
        cur_min = cur_max = 0
        if part in minmax_idx.index:
            if 'current_min' in minmax_idx.columns:
                cur_min = float(minmax_idx.at[part, 'current_min'])
            if 'current_max' in minmax_idx.columns:
                cur_max = float(minmax_idx.at[part, 'current_max'])

        # qty on hand / cost
        qty_oh = unit_cost = 0
        description = store = ''
        if part in stock_idx.index:
            if 'qty_on_hand' in stock_idx.columns:
                qty_oh = float(stock_idx.at[part, 'qty_on_hand'])
            if 'unit_cost' in stock_idx.columns:
                unit_cost = float(stock_idx.at[part, 'unit_cost'])
            if 'description' in stock_idx.columns:
                description = str(stock_idx.at[part, 'description'])
            if 'store' in stock_idx.columns:
                store = str(stock_idx.at[part, 'store'])

        # Recommended MIN / MAX
        if classification == 'Fast Mover':
            mean_daily = mean_monthly / 30
            safety_stock = 1.65 * std_monthly * math.sqrt(lead_days / 30)
            rec_min = math.ceil(mean_daily * lead_days + safety_stock)
            # order qty mode from REC-STOR transactions
            try:
                oq = mode_order_qty(goods_in_by_part.get_group(part))
            except KeyError:
                oq = max(10, round(mean_monthly * 1.5))
            rec_max = rec_min + oq

        elif classification == 'Slow Mover':
            rec_min = 1
            rec_max = 2

        else:  # Dead Stock
            rec_min = 0
            rec_max = 0

        # Flags
        overstocked        = qty_oh > 2 * rec_max if rec_max > 0 else False
        understocked       = qty_oh < rec_min
        min_variance       = rec_min - cur_min
        max_variance       = rec_max - cur_max

        part_stats.append({
            'part_number':       part,
            'description':       description,
            'store':             store,
            'classification':    classification,
            'qty_on_hand':       qty_oh,
            'unit_cost':         unit_cost,
            'stock_value':       round(qty_oh * unit_cost, 2),
            'mean_monthly_consumption': round(mean_monthly, 2),
            'std_monthly_consumption':  round(std_monthly, 2),
            'peak_monthly_consumption': round(peak_monthly, 2),
            'months_with_usage': months_with_usage,
            'total_12m_consumption': round(total_consumption, 2),
            'lead_time_days':    lead_days,
            'current_min':       cur_min,
            'current_max':       cur_max,
            'recommended_min':   rec_min,
            'recommended_max':   rec_max,
            'min_variance':      round(min_variance, 2),
            'max_variance':      round(max_variance, 2),
            'overstocked':       overstocked,
            'understocked':      understocked,
        })

    recs_df = pd.DataFrame(part_stats)

    # ── Exceptions ────────────────────────────────────────────────────────
    exc_rows = []
    for _, row in recs_df.iterrows():
        reasons = []
        if row['classification'] == 'Dead Stock':
            reasons.append('Dead Stock — no usage in 12 months')
        if row['overstocked']:
            reasons.append(f"Overstocked: {row['qty_on_hand']:.0f} on hand vs MAX {row['recommended_max']:.0f}")
        if row['understocked'] and row['classification'] != 'Dead Stock':
            reasons.append(f"Understocked: {row['qty_on_hand']:.0f} on hand vs MIN {row['recommended_min']:.0f}")
        if reasons:
            exc_rows.append({
                'part_number':    row['part_number'],
                'description':    row['description'],
                'store':          row['store'],
                'classification': row['classification'],
                'qty_on_hand':    row['qty_on_hand'],
                'recommended_min': row['recommended_min'],
                'recommended_max': row['recommended_max'],
                'exception_reason': '; '.join(reasons),
                'action_required': (
                    'Review for disposal' if row['classification'] == 'Dead Stock'
                    else 'Reduce order / return stock' if row['overstocked']
                    else 'Place urgent order'
                ),
            })
    exc_df = pd.DataFrame(exc_rows)

    # ── Monthly pivot (wide) ──────────────────────────────────────────────
    pivot = monthly_full.pivot_table(
        index='part_number', columns='month_str', values='consumption', aggfunc='sum', fill_value=0
    ).reset_index()

    # ── Summary stats ─────────────────────────────────────────────────────
    total_parts   = len(recs_df)
    fast_count    = int((recs_df['classification'] == 'Fast Mover').sum())
    slow_count    = int((recs_df['classification'] == 'Slow Mover').sum())
    dead_count    = int((recs_df['classification'] == 'Dead Stock').sum())
    over_count    = int(recs_df['overstocked'].sum())
    under_count   = int(recs_df['understocked'].sum())
    ok_count      = total_parts - over_count - under_count

    summary = {
        'total_parts':   total_parts,
        'fast_movers':   fast_count,
        'slow_movers':   slow_count,
        'dead_stock':    dead_count,
        'overstocked':   over_count,
        'understocked':  under_count,
        'correctly_stocked': ok_count,
        'total_stock_value': float(round(recs_df['stock_value'].sum(), 2)),
        'exceptions_count': len(exc_df),
    }

    # ── Chart data ────────────────────────────────────────────────────────
    # Top 20 overstocked
    top_overstock = (
        recs_df[recs_df['overstocked']]
        .nlargest(20, 'qty_on_hand')[['part_number', 'qty_on_hand', 'recommended_max']]
        .to_dict(orient='records')
    )

    # Monthly total consumption trend
    if all_months:
        monthly_trend = (
            monthly_full.groupby('month_str')['consumption']
            .sum()
            .reset_index()
            .sort_values('month_str')
        )
        trend_labels = monthly_trend['month_str'].tolist()
        trend_values = monthly_trend['consumption'].tolist()
    else:
        trend_labels = []
        trend_values = []

    chart_data = {
        'classification': {
            'labels': ['Fast Mover', 'Slow Mover', 'Dead Stock'],
            'values': [fast_count, slow_count, dead_count],
        },
        'stock_status': {
            'labels': ['Overstocked', 'Understocked', 'Correctly Stocked'],
            'values': [over_count, under_count, ok_count],
        },
        'top_overstock': top_overstock,
        'monthly_trend': {
            'labels': trend_labels,
            'values': [round(v, 1) for v in trend_values],
        },
    }

    return recs_df, exc_df, pivot, summary, chart_data


# ─────────────────────────────────────────────────────────────────────────────
# Excel output builder
# ─────────────────────────────────────────────────────────────────────────────

HEADER_FILL  = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT  = Font(color='FFFFFF', bold=True)
ALT_FILL     = PatternFill('solid', fgColor='D6E4F0')
WARN_FILL    = PatternFill('solid', fgColor='FFE0B2')
DANGER_FILL  = PatternFill('solid', fgColor='FFCDD2')
OK_FILL      = PatternFill('solid', fgColor='C8E6C9')
THIN_BORDER  = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)


def auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def write_sheet(ws, df, title_row=None):
    if df.empty:
        ws.append(['No data'])
        return

    headers = list(df.columns)

    if title_row:
        ws.append([title_row])
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=len(headers))
        title_cell = ws.cell(1, 1)
        title_cell.font = Font(bold=True, size=13, color='1F4E79')
        title_cell.alignment = Alignment(horizontal='center')
        data_start = 3
        ws.append([])
    else:
        data_start = 1

    # Header row
    ws.append(headers)
    header_row_idx = ws.max_row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(header_row_idx, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER

    # Data rows
    for row_idx, (_, row) in enumerate(df.iterrows()):
        ws.append([
            (v.item() if hasattr(v, 'item') else v)
            for v in row.values
        ])
        excel_row = ws.max_row
        fill = ALT_FILL if row_idx % 2 == 1 else None

        # Colour coding for recs sheet
        cls = row.get('classification', '')
        over = row.get('overstocked', False)
        under = row.get('understocked', False)
        if over:
            fill = WARN_FILL
        elif under:
            fill = DANGER_FILL

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(excel_row, col_idx)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal='left')
            if fill:
                cell.fill = fill

    auto_width(ws)


def build_excel(recs_df, exc_df, monthly_df):
    wb = Workbook()

    # Sheet 1: MRP Recommendations
    ws1 = wb.active
    ws1.title = 'MRP Recommendations'
    recs_display = recs_df[[
        'part_number', 'description', 'store', 'classification',
        'qty_on_hand', 'unit_cost', 'stock_value',
        'mean_monthly_consumption', 'std_monthly_consumption',
        'peak_monthly_consumption', 'months_with_usage', 'total_12m_consumption',
        'lead_time_days',
        'current_min', 'current_max',
        'recommended_min', 'recommended_max',
        'min_variance', 'max_variance',
        'overstocked', 'understocked',
    ]].copy()
    recs_display.columns = [
        'Part Number', 'Description', 'Store', 'Classification',
        'Qty On Hand', 'Unit Cost', 'Stock Value (£)',
        'Mean Monthly Consumption', 'Std Dev Monthly', 'Peak Monthly',
        'Months With Usage', 'Total 12M Consumption',
        'Lead Time (Days)',
        'Current MIN', 'Current MAX',
        'Recommended MIN', 'Recommended MAX',
        'MIN Variance', 'MAX Variance',
        'Overstocked?', 'Understocked?',
    ]
    write_sheet(ws1, recs_display, 'MRP Recommendations — ' + datetime.today().strftime('%d %b %Y'))

    # Sheet 2: Exceptions
    ws2 = wb.create_sheet('Exceptions')
    if not exc_df.empty:
        exc_display = exc_df.copy()
        exc_display.columns = [c.replace('_', ' ').title() for c in exc_display.columns]
        write_sheet(ws2, exc_display, 'Exceptions Requiring Review')
    else:
        ws2.append(['No exceptions found — all parts within acceptable parameters.'])

    # Sheet 3: Monthly Consumption
    ws3 = wb.create_sheet('Monthly Consumption')
    write_sheet(ws3, monthly_df, 'Monthly Consumption Breakdown (12 Months)')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ─────────────────────────────────────────────────────────────────────────────
# Flask routes
# ─────────────────────────────────────────────────────────────────────────────

# In-memory store for analysis results (single-user; extend with sessions if needed)
_analysis_cache = {}


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/health')
def health():
    return jsonify({'status': 'ok'})


@app.route('/detect-columns', methods=['POST'])
def detect_columns_route():
    """
    Called after file uploads to return detected column mappings.
    Client can then present them to the user for confirmation / correction.
    """
    results = {}
    try:
        for key, alias_map in [('stock', STOCK_ALIASES),
                                ('transactions', TRANS_ALIASES),
                                ('minmax', MINMAX_ALIASES)]:
            f = request.files.get(key)
            if not f:
                return jsonify({'error': f'Missing file: {key}'}), 400
            df = pd.read_excel(f, nrows=5)
            mapping = detect_columns(df, alias_map)
            results[key] = {
                'columns':    list(df.columns),
                'detected':   mapping,
                'required':   [k for k, v in mapping.items() if v is None],
            }
    except Exception as e:
        return jsonify({'error': str(e), 'trace': traceback.format_exc()}), 500

    return jsonify(results)


@app.route('/analyse', methods=['POST'])
def analyse():
    """
    Accepts files + optional column overrides, runs analysis,
    returns summary + chart data as JSON and caches Excel bytes.
    """
    try:
        stock_file = request.files.get('stock')
        trans_file = request.files.get('transactions')
        minmax_file = request.files.get('minmax')

        if not all([stock_file, trans_file, minmax_file]):
            return jsonify({'error': 'All three files are required.'}), 400

        stock_df  = pd.read_excel(stock_file)
        trans_df  = pd.read_excel(trans_file)
        minmax_df = pd.read_excel(minmax_file)

        # Parse optional user overrides from form fields
        def parse_overrides(prefix):
            overrides = {}
            for key in request.form:
                if key.startswith(prefix + '_'):
                    logical = key[len(prefix) + 1:]
                    val = request.form[key].strip()
                    if val:
                        overrides[logical] = val
            return overrides or None

        stock_overrides  = parse_overrides('stock')
        trans_overrides  = parse_overrides('trans')
        minmax_overrides = parse_overrides('minmax')

        stock_map  = detect_columns(stock_df,  STOCK_ALIASES)
        trans_map  = detect_columns(trans_df,  TRANS_ALIASES)
        minmax_map = detect_columns(minmax_df, MINMAX_ALIASES)

        recs_df, exc_df, monthly_df, summary, chart_data = run_mrp_analysis(
            stock_df, trans_df, minmax_df,
            stock_map, trans_map, minmax_map,
            stock_overrides, trans_overrides, minmax_overrides,
        )

        # Cache Excel
        excel_bytes = build_excel(recs_df, exc_df, monthly_df)
        _analysis_cache['excel'] = excel_bytes.read()
        _analysis_cache['filename'] = (
            'MRP_Analysis_' + datetime.today().strftime('%Y%m%d') + '.xlsx'
        )

        return jsonify({
            'summary':    summary,
            'chart_data': chart_data,
            'status':     'ok',
        })

    except Exception as e:
        return jsonify({'error': str(e), 'trace': traceback.format_exc()}), 500


@app.route('/download')
def download():
    if 'excel' not in _analysis_cache:
        return 'No analysis data available. Please run analysis first.', 404
    return send_file(
        io.BytesIO(_analysis_cache['excel']),
        download_name=_analysis_cache.get('filename', 'MRP_Analysis.xlsx'),
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.errorhandler(413)
def request_entity_too_large(_):
    return jsonify({'error': 'Upload too large. Maximum total upload size is 200 MB.'}), 413


if __name__ == '__main__':
    app.run(host='0.0.0.0', debug=True, port=5000)
